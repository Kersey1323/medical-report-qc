from tests.integration.test_schema import factory, seeded
from app.services.review_service import append_review
from app.services.export_service import request_export, build_export
from app.tasks.recovery import publish_outbox
from app.schemas.requests import ReviewCreate
from app.core.exceptions import AppError
from app.core.config import Settings
from app.db.models import ModuleResult, TaskOutbox
from sqlalchemy import select, func
from openpyxl import load_workbook
import pytest


def test_review_revision(factory, seeded):
    with factory.begin() as s:
        m = s.get(ModuleResult, seeded[2])
        m.status = "succeeded"
        m.finding = "issue"
        m.outcome = {
            "status": "succeeded",
            "finding": "issue",
            "suggestions": [{"text": "建议"}],
        }
    cmd = ReviewCreate(
        module_result_id=seeded[2], generation=1, type="defer", expected_revision=0
    )
    with factory.begin() as s:
        assert append_review(s, cmd)["revision"] == 1
    with pytest.raises(AppError):
        with factory.begin() as s:
            append_review(s, cmd)


def test_export_literal(factory, seeded, tmp_path):
    with factory.begin() as s:
        r = s.get(__import__("app.db.models", fromlist=["Report"]).Report, seeded[3])
        r.description = "=1+1"
        a = request_export(s, seeded[0])
        id = a["id"]
    path = build_export(id, factory, Settings(var_dir=tmp_path))
    w = load_workbook(path)
    assert w.sheetnames == [
        "报告汇总",
        "诊断质控",
        "字符纠错",
        "语法纠错",
        "片段定位",
        "失败记录",
        "处理汇总",
    ]
    matching = [c for row in w.worksheets[0] for c in row if c.value == "=1+1"]
    assert matching and matching[0].data_type == "s"


def test_outbox_failure_retained(factory, seeded):
    with factory.begin() as s:
        b = TaskOutbox(kind="qc", target_id=seeded[2], eligible=True)
        s.add(b)
        s.flush()
        id = b.id
        # This persistent test database can contain outboxes from earlier runs.
        # Include the new fixture in both passes instead of testing only the oldest 100.
        batch_size = s.scalar(select(func.count()).select_from(TaskOutbox))

    failed_targets = []

    def bad(*a):
        failed_targets.append(a)
        raise ConnectionError("broker down")

    publish_outbox(factory, publish=bad, limit=batch_size)
    assert ("qc", seeded[2]) in failed_targets
    with factory() as s:
        assert s.get(TaskOutbox, id).published_at is None
    delivered = []
    publish_outbox(factory, publish=lambda k, i: delivered.append((k, i)), limit=batch_size)
    assert ("qc", seeded[2]) in delivered
