"""Regression cases from the independent whole-branch review (synthetic data)."""

from datetime import timedelta
from types import SimpleNamespace
from uuid import UUID
import hashlib
import pytest
from sqlalchemy import select
from openpyxl import load_workbook
from tests.integration.test_schema import factory, seeded
from app.db.models import *
from app.core.config import Settings
from app.core.exceptions import AppError, StageValidationError
from contracts.errors import ErrorDetail
from contracts.inference import RequestState
from app.domain.results import StageContext, RunConfig, StageOutput, ReportInput
from app.services.execution_service import StageRunner


def uncertain_attempt(s, mid):
    a = InferenceAttempt(
        module_result_id=mid,
        generation=1,
        stage="csc.correct",
        chunk_key="c",
        attempt_no=1,
        input_hash="input",
        config_hash="x",
        remote_request_id=uid(),
        status="failed",
        error={
            "code": "upstream_timeout",
            "message": "uncertain",
            "stage": "csc.correct",
            "retryable": False,
            "remote_state": "unknown",
        },
    )
    s.add(a)
    return a


def test_failed_uncertain_gpu_is_reconciled_by_original_id(
    factory, seeded, monkeypatch
):
    from app.tasks.recovery import recover_once

    monkeypatch.setattr("app.tasks.recovery.publish_outbox", lambda *a, **kw: {})
    with factory.begin() as s:
        m = s.get(ModuleResult, seeded[2])
        m.status = "failed"
        a = uncertain_attempt(s, m.id)
        s.flush()
        rid = a.remote_request_id
    calls = []

    def lookup(id):
        calls.append(str(id))
        return RequestState(request_id=id, status="succeeded")

    recover_once(factory, lookup)
    assert calls.count(rid) == 1
    with factory() as s:
        assert s.get(ModuleResult, seeded[2]).status == "pending"
        a = s.scalar(
            select(InferenceAttempt).where(
                InferenceAttempt.module_result_id == seeded[2]
            )
        )
        assert a.status == "resumable" and a.error is None
    ctx = StageContext(UUID(seeded[2]), 1, RunConfig(fingerprint="x"), lambda: False)
    seen = []
    StageRunner(factory).run(
        "csc.correct",
        "c",
        "input",
        ctx,
        lambda id, n: seen.append(str(id)) or StageOutput(payload={"raw_text": "完成"}),
    )
    assert seen == [rid]


def test_retry_rejects_uncertainty_without_creating_generation(factory, seeded):
    from app.services.run_service import retry_failed

    with factory.begin() as s:
        m = s.get(ModuleResult, seeded[2])
        m.status = "failed"
        uncertain_attempt(s, m.id)
    with pytest.raises(AppError, match="未确认"):
        with factory.begin() as s:
            q = s.get(QualityRun, seeded[0])
            retry_failed(s, q.id, q.revision)
    with factory() as s:
        assert s.get(ModuleResult, seeded[2]).generation == 1


def test_resumed_validation_failure_uses_budget_and_preserves_audit(factory, seeded):
    from app.clients.inference import resuming_remote

    with factory.begin() as s:
        a = uncertain_attempt(s, seeded[2])
        a.status = "resumable"
        a.error = None
        s.flush()
        rid = a.remote_request_id
    calls = []

    def invoke(id, n):
        calls.append((str(id), n, resuming_remote.get()))
        if n == 1:
            raise StageValidationError(
                ErrorDetail(
                    code="response_invalid",
                    message="invalid",
                    stage="csc.correct",
                    retryable=True,
                ),
                "truncated-raw",
                {"model": "synthetic"},
            )
        return StageOutput(payload={"raw_text": "valid"})

    ctx = StageContext(UUID(seeded[2]), 1, RunConfig(fingerprint="x"), lambda: False)
    StageRunner(factory).run("csc.correct", "c", "input", ctx, invoke)
    assert (
        [c[1:] for c in calls] == [(1, True), (2, False)]
        and calls[0][0] == rid
        and calls[1][0] != rid
    )
    with factory() as s:
        a = s.scalar(
            select(InferenceAttempt).where(InferenceAttempt.remote_request_id == rid)
        )
        assert a.raw_response == "truncated-raw" and a.identity == {
            "model": "synthetic"
        }


@pytest.mark.parametrize("old_raises", [False, True])
def test_import_stale_worker_cannot_overwrite_success(
    factory, tmp_path, monkeypatch, old_raises
):
    import app.services.import_service as svc

    p = tmp_path / "synthetic.csv"
    p.write_text("描述\n合成文本\n")
    with factory.begin() as s:
        row = Import(filename=p.name, path=str(p), sha256="x")
        s.add(row)
        s.flush()
        id = row.id
    original = svc.parse_upload
    first = True

    def parse(*args):
        nonlocal first
        if first:
            first = False
            with factory.begin() as s:
                s.get(Import, id).lease_until = now() - timedelta(seconds=1)
            svc.process_import(id, factory)
            if old_raises:
                raise RuntimeError("old parse failed")
        return original(*args)

    monkeypatch.setattr(svc, "parse_upload", parse)
    svc.process_import(id, factory)
    with factory() as s:
        row = s.get(Import, id)
        assert row.status == "ready" and row.errors == []
        assert (
            len(s.scalars(select(ImportRow).where(ImportRow.import_id == id)).all())
            == 1
        )


@pytest.mark.parametrize("old_raises", [False, True])
def test_export_stale_worker_cannot_publish_or_fail_successor(
    factory, seeded, tmp_path, monkeypatch, old_raises
):
    import app.services.export_service as svc

    with factory.begin() as s:
        id = svc.request_export(s, seeded[0])["id"]
    settings = Settings(var_dir=tmp_path)
    save = svc.Workbook.save
    first = True
    temps = []
    winner = []

    def interleave(w, path):
        nonlocal first
        temps.append(str(path))
        if first:
            first = False
            with factory.begin() as s:
                s.get(Artifact, id).lease_until = now() - timedelta(seconds=1)
            winner.append(svc.build_export(id, factory, settings).read_bytes())
            if old_raises:
                raise RuntimeError("old export failed")
        return save(w, path)

    monkeypatch.setattr(svc.Workbook, "save", interleave)
    try:
        svc.build_export(id, factory, settings)
    except RuntimeError:
        if not old_raises:
            raise
    with factory() as s:
        row = s.get(Artifact, id)
        assert row.status == "ready" and row.error is None
        actual = __import__("pathlib").Path(row.path).read_bytes()
        assert actual == winner[0] and hashlib.sha256(actual).hexdigest() == row.sha256
    assert len(set(temps)) == 2


def test_export_preserves_extra_columns_failure_row_and_attempts(
    factory, seeded, tmp_path
):
    from app.services.export_service import request_export, build_export

    with factory.begin() as s:
        r = s.get(Report, seeded[3])
        r.extra_fields = {"外部编号": "=ID-42", "备注": "保留"}
        rr = s.get(RunReport, seeded[1])
        rr.original_row = 42
        m = s.get(ModuleResult, seeded[2])
        m.status = "failed"
        m.outcome = {
            "status": "failed",
            "finding": None,
            "error": {"stage": "csc.correct", "message": "合成失败"},
        }
        for n in range(1, 4):
            a = uncertain_attempt(s, m.id)
            a.attempt_no = n
            a.error = {"stage": "csc.correct", "message": "合成失败"}
        id = request_export(s, seeded[0])["id"]
    w = load_workbook(build_export(id, factory, Settings(var_dir=tmp_path)))
    rows = list(w["报告汇总"].values)
    assert "原始附加字段" in rows[0] and "=ID-42" in str(rows)
    failures = list(w["失败记录"].values)
    assert all(h in failures[0] for h in ["原始行号", "阶段", "尝试总数"])
    data = dict(zip(failures[0], failures[1]))
    assert (
        data["原始行号"] == "42"
        and data["尝试总数"] == "3"
        and data["阶段"] == "csc.correct"
    )


def test_job_lease_renews_and_stops_after_owner_changes(factory, tmp_path):
    from app.services.job_lease import RenewingLease
    import time

    with factory.begin() as s:
        row = Import(
            filename="synthetic.csv",
            path=str(tmp_path / "x"),
            sha256="x",
            status="running",
            lease_owner="original",
            lease_until=now() + timedelta(seconds=10),
        )
        s.add(row)
        s.flush()
        id = row.id
    with RenewingLease(factory, Import, id, "original", seconds=120, interval=0.03):
        deadline = time.monotonic() + 2
        while True:
            with factory() as s:
                renewed = s.get(Import, id).lease_until > now() + timedelta(seconds=60)
            if renewed:
                break
            assert time.monotonic() < deadline
            time.sleep(0.02)
        with factory.begin() as s:
            row = s.get(Import, id)
            row.lease_owner = "successor"
            row.lease_until = now() + timedelta(seconds=240)
        with factory() as s:
            expiry = s.get(Import, id).lease_until
        time.sleep(0.1)
        with factory() as s:
            row = s.get(Import, id)
            assert row.lease_owner == "successor" and row.lease_until == expiry


def test_export_retains_unmapped_invalid_original_cells(factory, seeded, tmp_path):
    from app.services.export_service import request_export, build_export

    with factory.begin() as s:
        imp = Import(
            filename="synthetic.csv", path="unused", sha256="x", status="ready"
        )
        s.add(imp)
        s.flush()
        row = ImportRow(
            import_id=imp.id,
            original_row=99,
            cells={"外部编号": "=INVALID-99", "描述": "", "诊断": "正常"},
            errors=[{"message": "空描述"}],
        )
        s.add(row)
        s.flush()
        s.add(
            RunReport(
                run_id=seeded[0],
                import_row_id=row.id,
                original_row=99,
                invalid=True,
                errors=row.errors,
            )
        )
        s.flush()
        id = request_export(s, seeded[0])["id"]
    w = load_workbook(build_export(id, factory, Settings(var_dir=tmp_path)))
    assert "=INVALID-99" in str(list(w["报告汇总"].values))
    assert "99" in str(list(w["失败记录"].values))
