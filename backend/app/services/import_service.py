import csv, zipfile, hashlib
from dataclasses import dataclass, field
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select, update
from app.db.models import Import, ImportRow, TaskOutbox, uid, now
from app.core.exceptions import fail, error_from_exception
from app.services.common import require
from app.services.report_service import create_report
from app.schemas.requests import ReportCreate
from datetime import timedelta


@dataclass
class ParsedRow:
    original_row: int
    cells: dict
    errors: list = field(default_factory=list)


@dataclass
class ParsedImport:
    columns: list
    rows: list
    errors: list = field(default_factory=list)
    sheet: str | None = None


def parse_upload(path: Path, media_type: str):
    if path.stat().st_size > 20 * 1024 * 1024:
        raise fail("input_validation_error", "文件超过20MiB", "import.parse")
    workbook = None
    try:
        if path.suffix.lower() == ".xlsx":
            with zipfile.ZipFile(path) as z:
                if sum(x.file_size for x in z.infolist()) > 200 * 1024 * 1024:
                    raise fail(
                        "input_validation_error", "XLSX解压超过200MiB", "import.parse"
                    )
                if any("vbaProject" in x.filename for x in z.infolist()):
                    raise fail(
                        "input_validation_error", "不接受含宏文件", "import.parse"
                    )
            workbook = load_workbook(path, read_only=True, data_only=False)
            sheet = workbook.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
            sheetname = sheet.title
            values = []
            for i, row in enumerate(rows):
                if i > 10000 or len(row) > 100:
                    raise fail(
                        "input_validation_error", "超过10000行或100列", "import.parse"
                    )
                values.append(list(row))
        else:
            sheetname = None
            with path.open(encoding="utf-8-sig", newline="") as f:
                values = []
                for i, row in enumerate(csv.reader(f, strict=True)):
                    if i > 10000 or len(row) > 100:
                        raise fail(
                            "input_validation_error",
                            "超过10000行或100列",
                            "import.parse",
                        )
                    values.append(row)
        if len(values) < 2:
            raise fail("input_validation_error", "文件为空或没有数据行", "import.parse")
        headers = [str(c).strip() if c is not None else "" for c in values[0]]
        if not headers or "" in headers or len(set(headers)) != len(headers):
            raise fail("input_validation_error", "表头为空或重复", "import.parse")
        parsed = []
        for i, row in enumerate(values[1:], 2):
            error = (
                []
                if len(row) <= len(headers)
                else [{"code": "extra_cells", "message": "数据列数超过表头"}]
            )
            parsed.append(
                ParsedRow(
                    i,
                    {
                        h: str(row[j]) if j < len(row) and row[j] is not None else ""
                        for j, h in enumerate(headers)
                    },
                    error,
                )
            )
        return ParsedImport(headers, parsed, sheet=sheetname)
    except (UnicodeError, csv.Error, zipfile.BadZipFile) as exc:
        raise fail(
            "input_validation_error", "文件编码或结构无效，CSV需UTF-8", "import.parse"
        ) from exc
    finally:
        if workbook:
            workbook.close()


def register_upload(s, filename, content, settings, request_id=None):
    suffix = Path(filename or "").suffix.lower()
    if suffix not in (".csv", ".xlsx") or len(content) > 20 * 1024 * 1024:
        raise fail("input_validation_error", "只接受20MiB以内CSV/XLSX", "import.upload")
    id = uid()
    directory = settings.var_dir / "uploads"
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    p = directory / (id + suffix)
    with p.open("xb") as f:
        f.write(content)
    p.chmod(0o600)
    row = Import(
        id=id,
        filename=Path(filename).name[:255],
        path=str(p),
        sha256=hashlib.sha256(content).hexdigest(),
        origin_request_id=request_id,
    )
    s.add(row)
    s.add(TaskOutbox(kind="import", target_id=id, eligible=True))
    s.flush()
    return row


def process_import(id, factory):
    from app.services.job_lease import RenewingLease

    owner = uid()
    lease = RenewingLease(factory, Import, id, owner)
    with factory.begin() as s:
        row = require(s, Import, id, True)
        if row.status in ("ready", "failed") or (
            row.lease_until and row.lease_until > now()
        ):
            return
        row.lease_until = now() + timedelta(seconds=120)
        row.status = "running"
        row.lease_owner = owner
        path = row.path
    with lease:
        try:
            parsed = parse_upload(Path(path), "")
            with factory.begin() as s:
                row = require(s, Import, id, True)
                if not lease.owns(row):
                    return
                if s.scalar(
                    select(ImportRow.id).where(ImportRow.import_id == id).limit(1)
                ):
                    raise fail("import_state_conflict", "导入行已存在", "import.parse")
                for x in parsed.rows:
                    s.add(
                        ImportRow(
                            import_id=id,
                            original_row=x.original_row,
                            cells=x.cells,
                            errors=x.errors,
                        )
                    )
                row.columns = parsed.columns
                row.status = "ready"
                row.lease_until = None
                row.lease_owner = None
        except Exception as exc:
            with factory.begin() as s:
                row = require(s, Import, id, True)
                if not lease.owns(row):
                    return
                row.status = "failed"
                row.errors = [error_from_exception(exc, "import.parse").model_dump()]
                row.lease_until = None
                row.lease_owner = None
            from app.core.logging_config import logger

            logger.opt(exception=exc).error("批量导入失败: task_id={}", id)


def confirm_import(s, id, mapping):
    row = require(s, Import, id, True)
    target = {
        "description_column": mapping.description_column,
        "diagnosis_column": mapping.diagnosis_column,
    }
    if row.confirmed and (row.mapping != target or row.checks != mapping.checks):
        raise fail(
            "import_mapping_conflict",
            "该导入已用其他映射或检查项确认",
            "import.confirm",
        )
    if row.status != "ready":
        raise fail("input_validation_error", "导入尚未解析成功", "import.confirm")
    if (
        mapping.description_column not in row.columns
        or (mapping.diagnosis_column and mapping.diagnosis_column not in row.columns)
        or ("diagnosis" in mapping.checks and not mapping.diagnosis_column)
    ):
        raise fail(
            "input_validation_error", "映射列不存在或诊断列未指定", "import.confirm"
        )
    rows = s.scalars(
        select(ImportRow)
        .where(ImportRow.import_id == id)
        .order_by(ImportRow.original_row)
    ).all()
    if not row.confirmed:
        for r in rows:
            description = r.cells.get(mapping.description_column, "")
            diagnosis = (
                r.cells.get(mapping.diagnosis_column, "")
                if mapping.diagnosis_column
                else ""
            )
            errors = list(r.errors)
            if not description.strip():
                errors.append({"code": "description_empty", "message": "描述为空"})
            if "diagnosis" in mapping.checks and not diagnosis.strip():
                errors.append({"code": "diagnosis_empty", "message": "诊断为空"})
            if len(description) > 200000 or len(diagnosis) > 200000:
                errors.append(
                    {"code": "text_too_long", "message": "文本超过200000字符"}
                )
            r.errors = errors
            if not errors:
                report = create_report(
                    s,
                    ReportCreate(
                        description=description,
                        diagnosis=diagnosis,
                        extra_fields={
                            k: v for k, v in r.cells.items() if k not in target.values()
                        },
                    ),
                    r.original_row,
                )
                r.report_id = report.id
        row.confirmed = True
        row.mapping = target
        row.checks = mapping.checks
    return {
        "report_ids": [r.report_id for r in rows if r.report_id],
        "invalid_rows": [
            {"original_row": r.original_row, "errors": r.errors}
            for r in rows
            if not r.report_id
        ],
    }
