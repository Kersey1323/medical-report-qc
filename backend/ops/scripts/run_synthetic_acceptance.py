"""Opt-in live synthetic checks. Never reads historical patient evaluation inputs."""

from pathlib import Path
import argparse, hashlib, json, os, time
import httpx


def run_acceptance(api, output, checks=None, batch=False, timeout=2400):
    checks = checks or ["diagnosis", "csc", "cgec", "rag"]
    out = Path(output)
    out.mkdir(parents=True, exist_ok=True, mode=0o700)
    manifest = {
        "status": "blocked",
        "submitted": 0,
        "checks": checks,
        "input_kind": "explicit_synthetic_fixture",
        "production_mock_calls": 0,
        "unverified": [
            "clinical quality",
            "full GPU boundary capacity",
            "worker process crash during active remote call",
        ],
        "errors": [],
    }

    def save():
        p = out / "manifest.json"
        p.write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
        p.chmod(0o600)

    if os.getenv("MEDICAL_QC_REAL_SMOKE") != "1":
        manifest["errors"] = ["Requires MEDICAL_QC_REAL_SMOKE=1"]
        save()
        return manifest
    start = time.monotonic()
    with httpx.Client(
        base_url=api.rstrip("/") + "/api/v1",
        timeout=180,
        transport=httpx.HTTPTransport(retries=0),
    ) as client:

        def call(method, path, **kw):
            r = client.request(method, path, **kw)
            r.raise_for_status()
            body = r.json()
            if body.get("status") != "success":
                raise RuntimeError(
                    json.dumps(body.get("detail", {}), ensure_ascii=False)
                )
            return body["data"]

        try:
            preflight = call("POST", "/system/preflight", json={"checks": checks})
            manifest["preflight"] = preflight
            if not preflight["ready"]:
                manifest["errors"] = ["Dependencies not ready"]
                save()
                return manifest
            fixture = json.loads(
                (
                    Path(__file__).resolve().parents[2]
                    / "tests/fixtures/synthetic_reports.json"
                ).read_text()
            )
            selected = fixture if batch else fixture[:1]
            ids = []
            hashes = {}
            for row in selected:
                r = call(
                    "POST",
                    "/reports",
                    json={k: row[k] for k in ["description", "diagnosis"]},
                )
                ids.append(r["id"])
                manifest["submitted"] += 1
                digest = hashlib.sha256(
                    json.dumps(
                        [row["description"], row["diagnosis"]], ensure_ascii=False
                    ).encode()
                ).hexdigest()
                hashes[r["id"]] = digest
                assert [r["description"], r["diagnosis"]] == [
                    row["description"],
                    row["diagnosis"],
                ]
            run = call("POST", "/runs", json={"report_ids": ids, "checks": checks})
            manifest["run_id"] = run["id"]
            manifest["original_report_sha256"] = hashes
            save()
            run = call("POST", f"/runs/{run['id']}/start")
            last = None
            while (
                run["status"] in ("pending", "running")
                and time.monotonic() - start < timeout
            ):
                counts = run["counts"]
                if counts != last:
                    print("Live synthetic progress:", json.dumps(counts), flush=True)
                    last = counts
                time.sleep(2)
                run = call("GET", f"/runs/{run['id']}")
            manifest["run_status"] = run["status"]
            manifest["counts"] = run["counts"]
            manifest["config_snapshot"] = run["config_snapshot"]
            rows = call("GET", f"/runs/{run['id']}/reports")["items"]
            results = []
            for row in rows:
                digest = hashlib.sha256(
                    json.dumps(
                        [row["description"], row["diagnosis"]], ensure_ascii=False
                    ).encode()
                ).hexdigest()
                assert hashes[row["report_id"]] == digest
                for branch, m in row["results"].items():
                    result = call("GET", f"/results/{m['id']}")
                    evidence = call("GET", f"/results/{m['id']}/evidence")
                    results.append(
                        {
                            "result_id": m["id"],
                            "branch": branch,
                            "generation": m["generation"],
                            "status": result["outcome"]["status"],
                            "finding": result["outcome"].get("finding"),
                            "error": result["outcome"].get("error"),
                            "attempts": [
                                {
                                    k: a[k]
                                    for k in [
                                        "stage",
                                        "attempt_no",
                                        "status",
                                        "identity",
                                    ]
                                }
                                for a in result["attempts"]
                            ],
                            "evidence_used": sum(e["used"] for e in evidence),
                            "full_reports": sum(
                                e["used"] and bool(e["candidate"]["full_report_text"])
                                for e in evidence
                            ),
                        }
                    )
            manifest["results"] = results
            artifact = call("POST", "/exports", json={"run_id": run["id"]})
            deadline = time.monotonic() + 90
            while (
                artifact["status"] in ("pending", "running")
                and time.monotonic() < deadline
            ):
                time.sleep(1)
                artifact = call("GET", f"/exports/{artifact['id']}")
            manifest["artifact_status"] = artifact["status"]
            if artifact["status"] == "ready":
                response = client.get(f"/exports/{artifact['id']}/download")
                assert "spreadsheetml" in response.headers.get("content-type", "")
                file = out / "actual.xlsx"
                file.write_bytes(response.content)
                file.chmod(0o600)
                from openpyxl import load_workbook

                w = load_workbook(file, read_only=True, data_only=False)
                assert len(w.sheetnames) == 7 and w.sheetnames[-1] == "处理汇总"
                manifest["xlsx_sheets"] = w.sheetnames
                w.close()
            manifest["status"] = (
                "passed"
                if all(r["status"] == "succeeded" for r in results)
                and len(results) == len(selected) * len(checks)
                and artifact["status"] == "ready"
                else "failed"
            )
        except Exception as exc:
            manifest["status"] = "failed"
            manifest["errors"].append(str(exc)[:1500])
    manifest["elapsed_seconds"] = round(time.monotonic() - start, 2)
    save()
    return manifest


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--api", default="http://127.0.0.1:5100")
    p.add_argument("--output", required=True)
    p.add_argument("--checks", nargs="+", choices=["diagnosis", "csc", "cgec", "rag"])
    p.add_argument("--batch", action="store_true")
    args = p.parse_args()
    result = run_acceptance(args.api, args.output, args.checks, args.batch)
    print(
        json.dumps(
            {
                k: result.get(k)
                for k in ["status", "run_id", "counts", "errors", "elapsed_seconds"]
            },
            ensure_ascii=False,
        )
    )
    raise SystemExit({"passed": 0, "blocked": 2}.get(result["status"], 1))
